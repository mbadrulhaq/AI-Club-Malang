#!/usr/bin/env python3
"""Tulis roster ke XLSX (multi-sheet) + CSV siap-import Fonnte."""
import json, re, csv
from collections import Counter
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter

ROWS = json.load(open('out/roster.json'))

# --- tandai kemungkinan akun bisnis / non-personal ---
BIZ = re.compile(r'\b(travel|property|properti|digital|sale|shop|shopee|store|toko|olshop|agency|'
                 r'studio|official|jasa|catering|klinik|villa|led|telkomsel|idcloud|zest|optima|'
                 r'admin|cs|marketing|assistant|iibf|com|net|id\b|group|corp|pt\b|cv\b)\b', re.I)
for r in ROWS:
    nm = r['nama_kontak'] or r['nama']
    r['catatan'] = 'Cek: kemungkinan akun bisnis' if (nm and BIZ.search(nm)) else ''

ORDER = {'Tinggi': 0, 'Sedang': 1, 'Rendah': 2, '-': 3}
ROWS.sort(key=lambda r: (r['status_grup'] != 'Masih di grup', r['gender'] != 'Perempuan',
                         -r['jml_pesan'], r['nama'] or 'zzz'))

COLS = [
    ('No', 6), ('Nama', 26), ('Nomor WA', 16), ('Gender (indikasi)', 17),
    ('Keyakinan Gender', 16), ('Dasar Gender', 24), ('Kota / Domisili', 18), ('Wilayah', 14),
    ('Sumber Nama', 19), ('Status Grup', 14), ('Tgl Gabung', 18), ('Jml Pesan', 10),
    ('Keaktifan', 14), ('Catatan', 26),
]
KEYS = ['_no', 'nama', 'wa', 'gender', 'keyakinan_gender', 'dasar_gender', 'kota', 'wilayah',
        'sumber_nama', 'status_grup', 'tgl_gabung', 'jml_pesan', 'keaktifan', 'catatan']

HDR_FILL = PatternFill('solid', fgColor='1F3864')
HDR_FONT = Font(name='Arial', size=10, bold=True, color='FFFFFF')
BODY = Font(name='Arial', size=10)
PINK = PatternFill('solid', fgColor='FCE4EC')
GREY = PatternFill('solid', fgColor='F2F2F2')


def sheet(wb, title, rows, note=None):
    ws = wb.create_sheet(title)
    top = 1
    if note:
        ws.cell(1, 1, note).font = Font(name='Arial', size=10, italic=True, color='555555')
        ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=len(COLS))
        top = 3
    for c, (h, w) in enumerate(COLS, 1):
        cell = ws.cell(top, c, h)
        cell.fill, cell.font = HDR_FILL, HDR_FONT
        cell.alignment = Alignment(vertical='center', wrap_text=True)
        ws.column_dimensions[get_column_letter(c)].width = w
    ws.row_dimensions[top].height = 30
    for i, r in enumerate(rows, 1):
        r = dict(r, _no=i)
        for c, k in enumerate(KEYS, 1):
            cell = ws.cell(top + i, c, r.get(k, ''))
            cell.font = BODY
            if k == 'wa' and r.get(k):
                cell.number_format = '@'
            if r['gender'] == 'Perempuan':
                cell.fill = PINK
            elif r['status_grup'] == 'Keluar':
                cell.fill = GREY
    ws.freeze_panes = ws.cell(top + 1, 1)
    ws.auto_filter.ref = '%s%d:%s%d' % ('A', top, get_column_letter(len(COLS)), top + len(rows))
    return ws


wb = Workbook()
wb.remove(wb.active)

aktif = [r for r in ROWS if r['status_grup'] == 'Masih di grup']
perempuan = [r for r in ROWS if r['gender'] == 'Perempuan']
bernama = [r for r in ROWS if r['nama']]
punya_wa = [r for r in aktif if r['wa']]
malang = [r for r in ROWS if r['wilayah'] == 'Malang Raya']

# --- Ringkasan ---
ws = wb.create_sheet('Ringkasan')
ws.column_dimensions['A'].width = 46
ws.column_dimensions['B'].width = 14
ws.column_dimensions['C'].width = 58
ws.cell(1, 1, 'AI CLUB MALANG — Hasil Ekstraksi Roster WhatsApp').font = Font(name='Arial', size=14, bold=True)
ws.cell(2, 1, 'Sumber: export chat grup "AI Club Malang 🔥" (8 Jul – 5 Sep 2026)').font = Font(name='Arial', size=10, italic=True, color='555555')
rows_sum = [
    ('METRIK', 'JUMLAH', 'CATATAN'),
    ('Total kontak unik terdeteksi', len(ROWS), 'Dari event join / added / left + pengirim pesan'),
    ('Masih di grup saat export', len(aktif), 'Ini basis untuk diundang ke grup baru'),
    ('Sudah keluar / dikeluarkan', len(ROWS) - len(aktif), 'Tetap disimpan, jangan dihapus'),
    ('', '', ''),
    ('Punya NOMOR WA (siap blasting)', sum(1 for r in ROWS if r['wa']), 'Langsung bisa diimport ke Fonnte'),
    ('  di antaranya masih di grup', len(punya_wa), 'Prioritas undangan gelombang 1'),
    ('Tanpa nomor (tersimpan di kontak HP Anda)', sum(1 for r in ROWS if not r['wa']), 'WhatsApp menampilkan nama, bukan nomor — ambil dari kontak HP'),
    ('', '', ''),
    ('Punya NAMA', len(bernama), 'Dari kontak HP + perkenalan diri di grup'),
    ('  dari kontak HP Anda', sum(1 for r in ROWS if r['nama_kontak']), ''),
    ('  dari perkenalan di grup', sum(1 for r in ROWS if not r['nama_kontak'] and r['nama_perkenalan']), 'Template "Nama: ... / Domisili: ..."'),
    ('Belum ketahuan namanya', len(ROWS) - len(bernama), 'Tidak pernah chat & tidak ada di kontak'),
    ('', '', ''),
    ('Ketahuan KOTA / DOMISILI', sum(1 for r in ROWS if r['kota']), ''),
    ('  Malang Raya', len(malang), 'Target utama kopdar offline'),
    ('  Luar Malang', sum(1 for r in ROWS if r['wilayah'] == 'Luar Malang'), 'Cocok untuk sesi online / kolaborasi lintas kota'),
    ('', '', ''),
    ('INDIKASI PEREMPUAN', len(perempuan), 'Kandidat grup khusus perempuan — WAJIB verifikasi manual'),
    ('Indikasi laki-laki', sum(1 for r in ROWS if r['gender'] == 'Laki-laki'), ''),
    ('Gender belum diketahui', sum(1 for r in ROWS if r['gender'] == 'Tidak diketahui'), 'Mayoritas tidak pernah chat — tanyakan lewat form pendaftaran'),
    ('', '', ''),
    ('Pernah kirim pesan', sum(1 for r in ROWS if r['jml_pesan']), ''),
    ('Aktif (>=10 pesan)', sum(1 for r in ROWS if r['jml_pesan'] >= 10), 'Kolam bibit pengurus & panitia'),
]
for i, (a, b, c) in enumerate(rows_sum, 4):
    ws.cell(i, 1, a).font = Font(name='Arial', size=10, bold=(i == 4 or a.isupper()))
    ws.cell(i, 2, b).font = Font(name='Arial', size=10, bold=(i == 4))
    ws.cell(i, 3, c).font = Font(name='Arial', size=10, color='555555')
    if i == 4:
        for c2 in range(1, 4):
            ws.cell(i, c2).fill, ws.cell(i, c2).font = HDR_FILL, HDR_FONT

sheet(wb, 'Semua Anggota', ROWS,
      'Seluruh kontak unik yang pernah terdeteksi di grup. Baris merah muda = indikasi perempuan; abu-abu = sudah keluar dari grup.')
sheet(wb, 'Undang ke Grup Baru', punya_wa,
      'Masih di grup saat export DAN punya nomor WA. Ini daftar utama untuk broadcast undangan grup baru. Kirim bertahap maks 100 nomor/jam lewat Fonnte.')
sheet(wb, 'Kandidat Grup Perempuan', perempuan,
      'Indikasi perempuan dari cara mereka menyebut diri di grup dan dari nama. WAJIB dicek satu per satu sebelum diundang — konfirmasi dulu lewat japri, jangan langsung dimasukkan ke grup.')
sheet(wb, 'Sudah Bernama', bernama,
      'Anggota yang namanya sudah diketahui — dari kontak HP Anda atau dari perkenalan diri di grup.')
sheet(wb, 'Malang Raya', malang,
      'Anggota yang menyebutkan domisili di Malang Raya. Prioritas undangan kopdar offline.')

wb.save('out/AICM_Roster_WhatsApp.xlsx')

# --- CSV ---
def dump(path, rows, cols):
    with open(path, 'w', newline='', encoding='utf-8') as f:
        w = csv.writer(f)
        w.writerow([c[0] for c in cols])
        for r in rows:
            w.writerow([r.get(c[1], '') for c in cols])

dump('out/fonnte_undang_grup_baru.csv', punya_wa,
     [('name', 'nama'), ('phone', 'wa'), ('kota', 'kota'), ('gender', 'gender')])
dump('out/fonnte_grup_perempuan.csv', [r for r in perempuan if r['wa']],
     [('name', 'nama'), ('phone', 'wa'), ('kota', 'kota'), ('keyakinan', 'keyakinan_gender'), ('dasar', 'dasar_gender')])
dump('out/roster_lengkap.csv', ROWS,
     [(h, k) for (h, _), k in zip(COLS[1:], KEYS[1:])])

print('OK')
print('  out/AICM_Roster_WhatsApp.xlsx  (6 sheet)')
print('  out/fonnte_undang_grup_baru.csv  ->', len(punya_wa), 'nomor')
print('  out/fonnte_grup_perempuan.csv    ->', len([r for r in perempuan if r['wa']]), 'nomor')
print('  out/roster_lengkap.csv           ->', len(ROWS), 'baris')
