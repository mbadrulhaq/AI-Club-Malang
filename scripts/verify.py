#!/usr/bin/env python3
"""Cek isi workbook hasil export + uji kualitas pencocokan dua grup."""
import json
from openpyxl import load_workbook

wb = load_workbook('out/AICM_Roster_WhatsApp.xlsx')
print('SHEET DALAM WORKBOOK')
for ws in wb:
    n = ws.max_row - 3 if ws.title != 'Ringkasan' else ws.max_row
    print('  %-26s %5d baris data x %2d kolom  filter=%s' % (
        ws.title, max(n, 0), ws.max_column, bool(ws.auto_filter.ref)))

lama = json.load(open('out/roster.json'))
baru = json.load(open('out/reborn.json'))
gab = json.load(open('out/roster.json'))  # hanya untuk hitung

ph_lama = {r['wa'] for r in lama if r['wa']}
nm_lama = {(r['nama_kontak'] or r['nama']).lower() for r in lama if not r['wa']}

only_name = [r for r in baru if not r['wa']]
cocok_nama = [r for r in only_name if (r['nama_kontak'] or r['nama']).lower() in nm_lama]
cocok_ph = [r for r in baru if r['wa'] and r['wa'] in ph_lama]

print()
print('KUALITAS PENCOCOKAN GRUP LAMA <-> REBORN')
print('  Anggota Reborn total          :', len(baru))
print('  - cocok lewat NOMOR WA        :', len(cocok_ph))
print('  - hanya nama (tanpa nomor)    :', len(only_name))
print('    di antaranya cocok nama lama:', len(cocok_nama))
print('    nama baru / tak cocok       :', len(only_name) - len(cocok_nama))
print('  => risiko hitung ganda maksimal: %d orang' % (len(only_name) - len(cocok_nama)))

print()
print('CONTOH — sheet "Belum Pindah ke Reborn" (5 teratas)')
ws = wb['Belum Pindah ke Reborn']
hdr = [c.value for c in ws[3]][:7]
print('  ' + ' | '.join(str(h)[:15] for h in hdr))
for row in ws.iter_rows(min_row=4, max_row=8, max_col=7, values_only=True):
    print('  ' + ' | '.join(str(v)[:15] if v is not None else '' for v in row))

print()
print('CONTOH — sheet "Kandidat Grup Perempuan"')
ws = wb['Kandidat Grup Perempuan']
for row in ws.iter_rows(min_row=4, max_row=ws.max_row, values_only=True):
    print('  %2s | %-20s | %-14s | %-18s | reborn=%-5s | %-14s | %s' % (
        row[0], str(row[1] or '-')[:20], str(row[2] or '(kontak HP)')[:14],
        str(row[3])[:18], row[5], str(row[9] or '-')[:14], str(row[8] or '')[:26]))
